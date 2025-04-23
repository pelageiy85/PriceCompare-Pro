import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import PatternFill
from openpyxl.utils.dataframe import dataframe_to_rows
from google.colab import files

# Установка библиотеки для работы с .xlsb
!pip install pyxlsb

# Загрузка файлов
print("Загрузите старый прайс...")
uploaded_old = files.upload()
old_file = list(uploaded_old.keys())[0]

print("Загрузите новый прайс...")
uploaded_new = files.upload()
new_file = list(uploaded_new.keys())[0]

# Чтение данных (формат .xlsb)
old_df = pd.read_excel(old_file, engine="pyxlsb")
new_df = pd.read_excel(new_file, engine="pyxlsb")

# Переименование колонок для избежания проблем с разделителями
old_df.columns = [str(col).replace(";", "_") for col in old_df.columns]
new_df.columns = [str(col).replace(";", "_") for col in new_df.columns]

# Объединение таблиц
merged_df = pd.merge(
    old_df,
    new_df,
    on=["Город отправления", "Город назначения"],
    suffixes=("_старый", "_новый"),
    how="outer"
)

# Список колонок с ценами (исключаем первые две колонки)
price_columns = [col for col in merged_df.columns if col not in ["Город отправления", "Город назначения"]]

# Расчет изменений
for col in price_columns:
    # Проверка, что колонка относится к старому или новому прайсу
    if col.endswith("_старый"):
        base_col = col.replace("_старый", "")
        new_col = f"{base_col}_новый"
        
        if new_col in merged_df.columns:
            merged_df[f"{base_col}_изменение"] = (
                (merged_df[new_col] - merged_df[col]) / 
                merged_df[col].replace(0, pd.NA) * 100
            ).round(2).fillna("N/A")

# Удаление временных колонок
result_df = merged_df.copy()
for col in price_columns:
    if col.endswith(("_старый", "_новый")):
        del result_df[col]



# Сохранение в Excel с форматированием
wb = Workbook()
ws = wb.active

# Запись данных
for r in dataframe_to_rows(result_df, index=False, header=True):
    ws.append(r)

# Условное форматирование для процентов
green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

for row in ws.iter_rows(min_row=2, min_col=3, max_col=ws.max_column):
    for cell in row:
        try:
            value = float(str(cell.value).replace("%", ""))
            if value < 0:
                cell.fill = green_fill
            elif value > 0:
                cell.fill = red_fill
        except (ValueError, AttributeError):
            pass

# Подсветка городов при наличии изменений [[3]][[7]]
for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
    # Проверяем, есть ли изменения в любой из колонок цен (столбцы C и далее)
    has_changes = any(
        cell.value not in ["N/A", None] and cell.value != 0 
        for cell in row[2:]  # Пропускаем первые 2 колонки (города)
    )
    if has_changes:
        # Закрашиваем ячейки с городами (A и B)
        ws.cell(row=row[0].row, column=1).fill = green_fill  # Город отправления
        ws.cell(row=row[0].row, column=2).fill = green_fill  # Город назначения

# Сохранение файла
output_file = "сравнительный_прайс.xlsx"
wb.save(output_file)
files.download(output_file)
print(f"Файл '{output_file}' готов к скачиванию.")
